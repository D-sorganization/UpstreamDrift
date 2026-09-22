"""Club-only matching UI façade for Motion Matching (CO-09 / #10613).

Thin service layer over workbook identity, matrix qualification, and fast-match
presets. GUI widgets bind here; solvers and stores stay in club_only / ledger.
"""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping

from src.shared.python.contracts import postcondition, precondition
from src.shared.python.motion_matching.club_only.fast_matching import MatchPreset
from src.shared.python.motion_matching.club_only.workbook_identity import (
    ALIAS_SHEETS,
    CANONICAL_TRIAL_SHEETS,
    sha256_file,
)

__all__ = [
    "BODY_MOTION_DISCLAIMER",
    "CLUB_ONLY_LANE",
    "ClubOnlyImportResult",
    "ClubOnlySessionState",
    "build_results_metadata",
    "import_club_only_workbook",
    "may_appear_as_verified",
    "observed_versus_inferred_legend",
    "verified_display_label",
]

BODY_MOTION_DISCLAIMER = (
    "Body motion is a plausible inferred candidate under golf priors — "
    "not measured body mocap."
)

CLUB_ONLY_LANE = "club_only"

_LEGEND: tuple[tuple[str, str], ...] = (
    ("observed_club", "Observed club (workbook measurement)"),
    (
        "inferred_body",
        "Inferred body candidate (plausible prior, not measured)",
    ),
)


@dataclass(frozen=True)
class ClubOnlyImportResult:
    """Outcome of importing a club-only Excel workbook for UI source selection."""

    workbook_path: Path
    trials: tuple[str, ...]
    missing_trials: tuple[str, ...]
    alias_conflicts: tuple[str, ...]
    coverage_notes: tuple[str, ...]
    ball_label_warnings: tuple[str, ...]
    errors: tuple[str, ...]
    ok: bool

    def as_dict(self) -> dict[str, Any]:
        return {
            "workbook_path": str(self.workbook_path),
            "trials": list(self.trials),
            "missing_trials": list(self.missing_trials),
            "alias_conflicts": list(self.alias_conflicts),
            "coverage_notes": list(self.coverage_notes),
            "ball_label_warnings": list(self.ball_label_warnings),
            "errors": list(self.errors),
            "ok": self.ok,
        }


@dataclass
class ClubOnlySessionState:
    """Editable club-only session; clone preserves user edits independently."""

    workbook_path: Path | None = None
    trial_id: str = ""
    model_id: str = ""
    preset: str = MatchPreset.FAST_PREVIEW.value
    user_notes: str = ""
    prior_edits: dict[str, Any] = field(default_factory=dict)
    cancel_requested: bool = False
    checkpoint_token: str | None = None

    def clone(self) -> ClubOnlySessionState:
        """Return an independent copy that preserves current edits."""
        return ClubOnlySessionState(
            workbook_path=self.workbook_path,
            trial_id=self.trial_id,
            model_id=self.model_id,
            preset=self.preset,
            user_notes=self.user_notes,
            prior_edits=deepcopy(self.prior_edits),
            cancel_requested=self.cancel_requested,
            checkpoint_token=self.checkpoint_token,
        )

    def request_cancel(self) -> None:
        self.cancel_requested = True

    def clear_cancel(self) -> None:
        self.cancel_requested = False

    def cancel_check(self) -> bool:
        """Fast-match cancel hook compatible with MatchCancelledError path."""
        return self.cancel_requested

    def with_checkpoint_token(self, token: str) -> ClubOnlySessionState:
        if not token:
            raise ValueError("checkpoint token must be non-empty")
        cloned = self.clone()
        cloned.checkpoint_token = token
        return cloned


def observed_versus_inferred_legend() -> tuple[tuple[str, str], ...]:
    """Stable legend entries for viewer and results panels."""
    return _LEGEND


@precondition(
    lambda matrix_cell_status: (
        isinstance(matrix_cell_status, str) and bool(matrix_cell_status.strip())
    ),
    "matrix_cell_status must be a non-empty string",
)
@postcondition(lambda result: isinstance(result, bool))
def may_appear_as_verified(*, matrix_cell_status: str) -> bool:
    """Only CO-08 matrix cells with status==scored may appear verified."""
    return matrix_cell_status.strip().lower() == "scored"


@precondition(
    lambda matrix_cell_status, preset: (
        isinstance(matrix_cell_status, str)
        and isinstance(preset, str)
        and bool(matrix_cell_status.strip())
        and bool(preset.strip())
    ),
    "matrix_cell_status and preset must be non-empty strings",
)
@postcondition(lambda result: isinstance(result, str) and bool(result))
def verified_display_label(*, matrix_cell_status: str, preset: str) -> str:
    """UI badge text; unqualified/fixture/preview never read as VERIFIED."""
    status = matrix_cell_status.strip().lower()
    preset_key = preset.strip().lower()
    if preset_key == MatchPreset.FAST_PREVIEW.value:
        return "PREVIEW"
    if status == "scored" and preset_key == MatchPreset.VERIFIED_FIT.value:
        return "VERIFIED"
    if status == "unqualified":
        return "UNQUALIFIED"
    if status == "rejected":
        return "REJECTED"
    if status == "unsupported":
        return "UNSUPPORTED"
    if status == "missing_runtime":
        return "MISSING RUNTIME"
    return "NOT VERIFIED"


@precondition(
    lambda trial_id, model_id, matrix_cell_status, preset, workbook_sha256: (
        bool(str(trial_id).strip())
        and bool(str(model_id).strip())
        and bool(str(matrix_cell_status).strip())
        and bool(str(preset).strip())
        and isinstance(workbook_sha256, str)
        and len(workbook_sha256) == 64
    ),
    "trial/model/status/preset required; workbook_sha256 must be 64-char hex",
)
@postcondition(
    lambda result: (
        result.get("lane") == CLUB_ONLY_LANE
        and result.get("source_kind") == "club_only_excel"
        and result.get("storage") == "matched_swing_ledger"
    )
)
def build_results_metadata(
    *,
    trial_id: str,
    model_id: str,
    matrix_cell_status: str,
    preset: str,
    workbook_sha256: str,
    extra: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Metadata for the existing matched-swing ledger — no parallel store."""
    appears = (
        may_appear_as_verified(matrix_cell_status=matrix_cell_status)
        and preset.strip().lower() == MatchPreset.VERIFIED_FIT.value
    )
    payload: dict[str, Any] = {
        "lane": CLUB_ONLY_LANE,
        "source_kind": "club_only_excel",
        "storage": "matched_swing_ledger",
        "trial_id": trial_id,
        "model_id": model_id,
        "matrix_cell_status": matrix_cell_status.strip().lower(),
        "preset": preset.strip().lower(),
        "workbook_sha256": workbook_sha256.lower(),
        "body_motion": "plausible_inferred_candidate",
        "body_motion_disclaimer": BODY_MOTION_DISCLAIMER,
        "appears_verified": appears,
        "legend": dict(_LEGEND),
    }
    if extra:
        for key, value in extra.items():
            if key in payload:
                raise ValueError(f"extra metadata cannot override reserved key {key!r}")
            payload[key] = value
    return payload


def _sheet_names(path: Path) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _open_workbook_errors() -> tuple[type[BaseException], ...]:
    errors: list[type[BaseException]] = [OSError, ValueError, KeyError, ImportError]
    try:
        from openpyxl.utils.exceptions import InvalidFileException

        errors.append(InvalidFileException)
    except ImportError:
        pass
    return tuple(errors)


@precondition(
    lambda path: path is not None and str(path).strip() != "",
    "workbook path is required",
)
@postcondition(lambda result: isinstance(result, ClubOnlyImportResult))
def import_club_only_workbook(path: Path | str) -> ClubOnlyImportResult:
    """List unique canonical trials, alias conflicts, and coverage for UI."""
    target = Path(path)
    if not target.is_file():
        return ClubOnlyImportResult(
            workbook_path=target,
            trials=(),
            missing_trials=CANONICAL_TRIAL_SHEETS,
            alias_conflicts=(),
            coverage_notes=(),
            ball_label_warnings=(),
            errors=(f"workbook path does not exist: {target}",),
            ok=False,
        )

    try:
        names = _sheet_names(target)
    except _open_workbook_errors() as exc:
        return ClubOnlyImportResult(
            workbook_path=target,
            trials=(),
            missing_trials=CANONICAL_TRIAL_SHEETS,
            alias_conflicts=(),
            coverage_notes=(),
            ball_label_warnings=(),
            errors=(f"failed to open workbook: {exc}",),
            ok=False,
        )

    present = set(names)
    trials = tuple(sheet for sheet in CANONICAL_TRIAL_SHEETS if sheet in present)
    missing = tuple(sheet for sheet in CANONICAL_TRIAL_SHEETS if sheet not in present)

    alias_conflicts: list[str] = []
    for alias in sorted(ALIAS_SHEETS):
        if alias not in present:
            continue
        # Filtering Experiments is the known alias of TW_ProV1 (CO-00).
        if "TW_ProV1" in present:
            alias_conflicts.append(
                f"{alias} aliases TW_ProV1 (deduplicated; not a fifth trial)"
            )
        else:
            alias_conflicts.append(f"{alias} present without TW_ProV1 primary sheet")

    coverage_notes: list[str] = []
    if trials:
        coverage_notes.append(
            f"{len(trials)}/{len(CANONICAL_TRIAL_SHEETS)} canonical trials present"
        )
    coverage_notes.append(
        "Observation coverage uses club markers only; body is not measured"
    )

    ball_label_warnings: list[str] = []
    errors: list[str] = []
    if not trials:
        errors.append(
            "no canonical trial sheets found "
            f"(expected one of {', '.join(CANONICAL_TRIAL_SHEETS)})"
        )

    # Hash for provenance when import succeeds enough to list trials.
    if trials:
        try:
            sha256_file(target)
        except ValueError as exc:
            errors.append(str(exc))

    ok = not errors and not missing
    if missing and trials:
        # Partial import is usable for the present trials but not fully ok.
        coverage_notes.append(f"missing trials: {', '.join(missing)}")
        ok = False
        errors.append(f"missing canonical trials: {', '.join(missing)}")

    return ClubOnlyImportResult(
        workbook_path=target.resolve(),
        trials=trials,
        missing_trials=missing,
        alias_conflicts=tuple(alias_conflicts),
        coverage_notes=tuple(coverage_notes),
        ball_label_warnings=tuple(ball_label_warnings),
        errors=tuple(errors),
        ok=ok,
    )
