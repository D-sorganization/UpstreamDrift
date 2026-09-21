"""Club-only workbook identity, units, events and trial lineage (CO-00 #10604).

Freezes content-addressed workbook manifests and the reviewed unit/frame/event
decisions. Does not alter source workbooks. Downstream loaders consume these
contracts; physical matching remains out of scope.
"""

from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping

from src.shared.python.contracts import postcondition, precondition
from src.shared.python.motion_matching.loaders.event_labels import (
    parse_event_marker_cells,
)

IDENTITY_SCHEMA = "club-workbook-identity/1.0.0"
NATIVE_SAMPLE_RATE_HZ = 240.0

# Content hashes from the 2026-09-20 club/neural review (PR #10628).
CLUB_DATA_SHA256 = "5d9183e1d01ea7c6f9c162375dd6855c076ee26a96b76c90e59e9cf2679dde25"
WIFFLE_PROV1_SHA256 = "88d3eb31541d886031f6c5ad7c82493b0e3ee4e3f73a8652372277cc14d02234"

CLUB_DATA_RELATIVE = Path("data/Club_Data.xlsx")
WIFFLE_PROV1_RELATIVE = Path(
    "src/engines/Simscape_Multibody_Models/3D_Golf_Model/matlab/src/apps/"
    "golf_gui/Motion Capture Plotter/Wiffle_ProV1_club_3D_data.xlsx"
)

# Canonical four trials; Filtering Experiments is an alias of TW_ProV1.
CANONICAL_TRIAL_SHEETS: tuple[str, ...] = (
    "TW_wiffle",
    "TW_ProV1",
    "GW_wiffle",
    "GW_ProV11",
)
ALIAS_SHEETS: frozenset[str] = frozenset({"Filtering Experiments"})

# Reviewed sample counts exclude trailing blank worksheet padding (max_row 885).
EXPECTED_SAMPLE_COUNTS: Mapping[str, int] = {
    "TW_wiffle": 882,
    "TW_ProV1": 775,
    "GW_wiffle": 774,
    "GW_ProV11": 771,
    "Filtering Experiments": 775,
}

EXPECTED_EVENT_SAMPLES: Mapping[str, Mapping[str, float]] = {
    "TW_wiffle": {"A": 240.0, "T": 412.0, "I": 519.0, "F": 832.0, "CHS": 114.5},
    "TW_ProV1": {"A": 240.0, "T": 418.0, "I": 525.0, "F": 725.0, "CHS": 114.5},
    "GW_wiffle": {"A": 240.0, "T": 448.0, "I": 517.0, "F": 724.0, "CHS": 104.6},
    "GW_ProV11": {"A": 240.0, "T": 452.0, "I": 521.0, "F": 721.0, "CHS": 115.1},
    "Filtering Experiments": {
        "A": 240.0,
        "T": 418.0,
        "I": 525.0,
        "F": 725.0,
        "CHS": 114.5,
    },
}


@dataclass(frozen=True)
class UnitAuthority:
    """Declared workbook units versus the reviewed SI interpretation."""

    declared_units: str
    reviewed_units: str
    to_meters_scale: float
    authority: str
    inches_scale_rejected: float

    def as_dict(self) -> dict[str, Any]:
        return {
            "declared_units": self.declared_units,
            "reviewed_units": self.reviewed_units,
            "to_meters_scale": self.to_meters_scale,
            "authority": self.authority,
            "inches_scale_rejected": self.inches_scale_rejected,
        }


@dataclass(frozen=True)
class FrameAuthority:
    """Recorded global/local frame semantics and known prose conflicts."""

    definitions_global_x: str
    definitions_global_y: str
    definitions_global_z: str
    definitions_local_z: str
    prose_conflict: str
    status: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "definitions_global_x": self.definitions_global_x,
            "definitions_global_y": self.definitions_global_y,
            "definitions_global_z": self.definitions_global_z,
            "definitions_local_z": self.definitions_local_z,
            "prose_conflict": self.prose_conflict,
            "status": self.status,
        }


@dataclass(frozen=True)
class EventAuthority:
    """Native clock and event-label semantics for club-only sheets."""

    sample_rate_hz: float
    impact_time_s: float
    clock: str
    t_event_meaning: str
    label_normalization: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "sample_rate_hz": self.sample_rate_hz,
            "impact_time_s": self.impact_time_s,
            "clock": self.clock,
            "t_event_meaning": self.t_event_meaning,
            "label_normalization": self.label_normalization,
        }


@dataclass(frozen=True)
class OrientationAxisPolicy:
    """Third-axis derivation policy for missing direction-cosine columns."""

    measured_axes: tuple[str, ...]
    missing_columns: tuple[str, ...]
    derived_axis: str
    status_when_derived: str
    degeneracy_flag_required: bool

    def as_dict(self) -> dict[str, Any]:
        return {
            "measured_axes": list(self.measured_axes),
            "missing_columns": list(self.missing_columns),
            "derived_axis": self.derived_axis,
            "status_when_derived": self.status_when_derived,
            "degeneracy_flag_required": self.degeneracy_flag_required,
        }


UNIT_AUTHORITY = UnitAuthority(
    declared_units="inches",
    reviewed_units="centimetres",
    to_meters_scale=0.01,
    authority=(
        "Definitions sheet declares inches; repository MATLAB/Python loaders "
        "and grip-to-face length (~1.07 m) qualify centimetre interpretation."
    ),
    inches_scale_rejected=0.0254,
)

FRAME_AUTHORITY = FrameAuthority(
    definitions_global_x="positive away from the target",
    definitions_global_y="positive toward the ball",
    definitions_global_z="positive up",
    definitions_local_z="positive from club head toward grip",
    prose_conflict=(
        "docs/motion_training/README.md describes X toward target and local X "
        "along shaft; Definitions sheet is retained as workbook authority until "
        "an explicit rigid transform is selected."
    ),
    status="conflict_recorded",
)

EVENT_AUTHORITY = EventAuthority(
    sample_rate_hz=NATIVE_SAMPLE_RATE_HZ,
    impact_time_s=0.0,
    clock="native impact-relative timestamps at 240 Hz",
    t_event_meaning=(
        "Pelvis-rotation event as documented; no body trajectory is present — "
        "do not relabel T as club top."
    ),
    label_normalization="strip trailing '=' from A/T/I/F/CHS labels",
)

ORIENTATION_AXIS_POLICY = OrientationAxisPolicy(
    measured_axes=("X", "Y"),
    missing_columns=("Zx", "Zy", "Zz"),
    derived_axis="Z",
    status_when_derived="derived_not_measured",
    degeneracy_flag_required=True,
)


@dataclass(frozen=True)
class WorkbookManifest:
    """Hash-verified workbook identity."""

    workbook_id: str
    relative_path: str
    sha256: str
    verified: bool

    def as_dict(self) -> dict[str, Any]:
        return {
            "workbook_id": self.workbook_id,
            "relative_path": self.relative_path,
            "sha256": self.sha256,
            "verified": self.verified,
        }


@dataclass(frozen=True)
class TrialRecord:
    """One unique club-only trial with lineage and native events."""

    trial_id: str
    primary_sheet: str
    alias_sheets: tuple[str, ...]
    lineage_id: str
    numeric_sample_count: int
    event_samples: Mapping[str, float]
    sheet_a1_label: str
    ball_label: str | None
    ball_label_status: str
    time_range_s: tuple[float, float]

    def as_dict(self) -> dict[str, Any]:
        return {
            "trial_id": self.trial_id,
            "primary_sheet": self.primary_sheet,
            "alias_sheets": list(self.alias_sheets),
            "lineage_id": self.lineage_id,
            "numeric_sample_count": self.numeric_sample_count,
            "event_samples": dict(self.event_samples),
            "sheet_a1_label": self.sheet_a1_label,
            "ball_label": self.ball_label,
            "ball_label_status": self.ball_label_status,
            "time_range_s": list(self.time_range_s),
        }


@dataclass(frozen=True)
class ClubWorkbookIdentity:
    """Frozen CO-00 identity package for both club workbooks."""

    schema: str
    manifests: tuple[WorkbookManifest, ...]
    trials: tuple[TrialRecord, ...]
    units: UnitAuthority
    frames: FrameAuthority
    events: EventAuthority
    orientation: OrientationAxisPolicy

    def as_dict(self) -> dict[str, Any]:
        return {
            "schema": self.schema,
            "manifests": [m.as_dict() for m in self.manifests],
            "trials": [t.as_dict() for t in self.trials],
            "units": self.units.as_dict(),
            "frames": self.frames.as_dict(),
            "events": self.events.as_dict(),
            "orientation": self.orientation.as_dict(),
        }

    def write_json(self, path: Path | str) -> None:
        target = Path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("w", encoding="utf-8") as handle:
            json.dump(self.as_dict(), handle, indent=2, allow_nan=False)
            handle.write("\n")


def sha256_file(path: Path | str) -> str:
    """Return lowercase hex SHA-256 of file bytes."""
    target = Path(path)
    if not target.is_file():
        raise ValueError(f"workbook path does not exist: {target}")
    digest = hashlib.sha256()
    with target.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def verify_workbook_hash(path: Path | str, expected_sha256: str) -> str:
    """Verify workbook content hash; fail closed on mismatch."""
    if not expected_sha256 or len(expected_sha256) != 64:
        raise ValueError("expected_sha256 must be a 64-char hex digest")
    actual = sha256_file(path)
    if actual != expected_sha256.lower():
        raise ValueError(
            f"workbook hash mismatch for {Path(path)}: "
            f"expected {expected_sha256}, got {actual}"
        )
    return actual


def _require_openpyxl() -> Any:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - env dependent
        raise ImportError("openpyxl is required for club workbook identity") from exc
    return load_workbook


def _open_workbook_sheet(path: Path | str, sheet_name: str) -> tuple[Any, Any]:
    """Open a workbook sheet after path/name validation; caller must close workbook."""
    target = Path(path)
    if not target.is_file():
        raise ValueError(f"workbook path does not exist: {target}")
    if not sheet_name or not str(sheet_name).strip():
        raise ValueError("sheet_name must be non-blank")
    load_workbook = _require_openpyxl()
    workbook = load_workbook(target, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"sheet {sheet_name!r} not found in {target.name}")
    return workbook, workbook[sheet_name]


def count_numeric_samples(path: Path | str, sheet_name: str) -> int:
    """Count non-blank numeric sample rows (column Sample #), ignoring padding."""
    workbook, sheet = _open_workbook_sheet(path, sheet_name)
    try:
        count = 0
        for row in sheet.iter_rows(min_row=4, max_col=1, values_only=True):
            value = row[0]
            if value is None:
                continue
            try:
                float(value)
            except (TypeError, ValueError):
                continue
            count += 1
        return count
    finally:
        workbook.close()


def read_sheet_event_samples(path: Path | str, sheet_name: str) -> dict[str, float]:
    """Read native event sample IDs from row 1 using shared label normalization."""
    workbook, sheet = _open_workbook_sheet(path, sheet_name)
    try:
        row1 = next(sheet.iter_rows(min_row=1, max_row=1, max_col=26, values_only=True))
        return parse_event_marker_cells(list(row1))
    finally:
        workbook.close()


def read_sheet_a1_label(path: Path | str, sheet_name: str) -> str:
    """Return the A1 ball/source label cell as a string (may be conflicting)."""
    workbook, sheet = _open_workbook_sheet(path, sheet_name)
    try:
        value = sheet["A1"].value
        if value is None:
            return ""
        return str(value).strip()
    finally:
        workbook.close()


def read_native_time_range(path: Path | str, sheet_name: str) -> tuple[float, float]:
    """Return (first, last) native impact-relative times for numeric samples."""
    workbook, sheet = _open_workbook_sheet(path, sheet_name)
    try:
        first: float | None = None
        last: float | None = None
        for row in sheet.iter_rows(min_row=4, max_col=2, values_only=True):
            sample, time_s = row[0], row[1]
            if sample is None or time_s is None:
                continue
            try:
                float(sample)
                t_val = float(time_s)
            except (TypeError, ValueError):
                continue
            if first is None:
                first = t_val
            last = t_val
        if first is None or last is None:
            raise ValueError(f"no numeric time samples on sheet {sheet_name!r}")
        return first, last
    finally:
        workbook.close()


def sheet_numeric_fingerprint(path: Path | str, sheet_name: str) -> str:
    """SHA-256 of Sample#/Time/position columns for lineage alias detection."""
    workbook, sheet = _open_workbook_sheet(path, sheet_name)
    try:
        digest = hashlib.sha256()
        for row in sheet.iter_rows(min_row=4, max_col=5, values_only=True):
            sample = row[0]
            if sample is None:
                continue
            try:
                float(sample)
            except (TypeError, ValueError):
                continue
            payload = "|".join("" if v is None else repr(float(v)) for v in row[:5])
            digest.update(payload.encode("utf-8"))
            digest.update(b"\n")
        return digest.hexdigest()
    finally:
        workbook.close()


def _ball_label_for_sheet(sheet_name: str, a1_label: str) -> tuple[str | None, str]:
    """Return (ball_label, status) without inventing missing labels."""
    if sheet_name == "GW_wiffle":
        # Sheet name says wiffle; A1 says ProV1 — retain the conflict.
        return None, "conflict_sheet_name_vs_a1"
    if a1_label.lower().startswith("wiffle"):
        return "wiffle", "stated_in_a1"
    if "prov1" in a1_label.lower().replace(" ", ""):
        return "ProV1", "stated_in_a1"
    if not a1_label:
        return None, "unknown"
    return None, "unknown"


@precondition(
    lambda repo_root: Path(repo_root).is_dir(),
    "repo_root must be an existing directory",
)
@postcondition(
    lambda result: result.schema == IDENTITY_SCHEMA,
    "identity package must use club-workbook-identity/1.0.0",
)
def build_club_workbook_identity(repo_root: Path | str) -> ClubWorkbookIdentity:
    """Build and verify the frozen club-only workbook identity package."""
    root = Path(repo_root)
    club_path = root / CLUB_DATA_RELATIVE
    wiffle_path = root / WIFFLE_PROV1_RELATIVE

    club_hash = verify_workbook_hash(club_path, CLUB_DATA_SHA256)
    wiffle_hash = verify_workbook_hash(wiffle_path, WIFFLE_PROV1_SHA256)

    manifests = (
        WorkbookManifest(
            workbook_id="club_data",
            relative_path=CLUB_DATA_RELATIVE.as_posix(),
            sha256=club_hash,
            verified=True,
        ),
        WorkbookManifest(
            workbook_id="wiffle_prov1_club_3d",
            relative_path=WIFFLE_PROV1_RELATIVE.as_posix(),
            sha256=wiffle_hash,
            verified=True,
        ),
    )

    fingerprints = {
        sheet: sheet_numeric_fingerprint(club_path, sheet)
        for sheet in (*CANONICAL_TRIAL_SHEETS, *sorted(ALIAS_SHEETS))
    }
    if fingerprints["Filtering Experiments"] != fingerprints["TW_ProV1"]:
        raise ValueError(
            "Filtering Experiments must share numeric lineage with TW_ProV1"
        )

    trials: list[TrialRecord] = []
    for sheet in CANONICAL_TRIAL_SHEETS:
        count = count_numeric_samples(club_path, sheet)
        expected = EXPECTED_SAMPLE_COUNTS[sheet]
        if count != expected:
            raise ValueError(
                f"sample count for {sheet}: expected {expected}, got {count}"
            )
        # Worksheet padding must not be counted as samples.
        if count >= 885:
            raise ValueError(
                f"sample count for {sheet} must exclude trailing blanks; got {count}"
            )
        events = read_sheet_event_samples(club_path, sheet)
        expected_events = EXPECTED_EVENT_SAMPLES[sheet]
        for key, value in expected_events.items():
            if key not in events or events[key] != value:
                raise ValueError(
                    f"event mismatch on {sheet} for {key}: "
                    f"expected {value}, got {events.get(key)}"
                )
        a1 = read_sheet_a1_label(club_path, sheet)
        ball_label, ball_status = _ball_label_for_sheet(sheet, a1)
        time_range = read_native_time_range(club_path, sheet)
        aliases: tuple[str, ...] = ()
        if sheet == "TW_ProV1":
            aliases = ("Filtering Experiments",)
        trials.append(
            TrialRecord(
                trial_id=sheet,
                primary_sheet=sheet,
                alias_sheets=aliases,
                lineage_id=fingerprints[sheet],
                numeric_sample_count=count,
                event_samples=events,
                sheet_a1_label=a1,
                ball_label=ball_label,
                ball_label_status=ball_status,
                time_range_s=time_range,
            )
        )

    # Cross-workbook shared-trial sheets must match lineage on overlapping names.
    for sheet in CANONICAL_TRIAL_SHEETS:
        other_fp = sheet_numeric_fingerprint(wiffle_path, sheet)
        if other_fp != fingerprints[sheet]:
            raise ValueError(
                f"cross-workbook lineage mismatch for sheet {sheet}: "
                f"Club_Data vs Wiffle_ProV1"
            )

    return ClubWorkbookIdentity(
        schema=IDENTITY_SCHEMA,
        manifests=manifests,
        trials=tuple(trials),
        units=UNIT_AUTHORITY,
        frames=FRAME_AUTHORITY,
        events=EVENT_AUTHORITY,
        orientation=ORIENTATION_AXIS_POLICY,
    )
