# Copyright (c) 2026 D-Sorganization. All rights reserved.
"""Tests for export module -- GIF, PNG, PDF export."""

from __future__ import annotations

import matplotlib
import numpy as np
import pytest
from matplotlib.figure import Figure

matplotlib.use("Agg")

from movement_optimizer.export import (
    _validate_export_path,
    export_animation_gif,
    export_plots_pdf,
    export_plots_png,
)


def _make_figure() -> Figure:
    """Create a simple matplotlib figure for testing."""
    fig = Figure(figsize=(4, 3))
    ax = fig.add_subplot(111)
    ax.plot([0, 1, 2], [0, 1, 0])
    return fig


class TestExportPNG:
    def test_produces_nonempty_file(self, tmp_path):
        fig = _make_figure()
        path = tmp_path / "test.png"
        export_plots_png(fig, str(path))
        assert path.exists()
        assert path.stat().st_size > 0

    def test_creates_parent_dirs(self, tmp_path):
        fig = _make_figure()
        path = tmp_path / "sub" / "dir" / "test.png"
        export_plots_png(fig, str(path))
        assert path.exists()


class TestExportPDF:
    def test_produces_nonempty_file(self, tmp_path):
        fig = _make_figure()
        path = tmp_path / "test.pdf"
        export_plots_pdf(fig, str(path))
        assert path.exists()
        assert path.stat().st_size > 0


class TestExportGIF:
    def test_produces_nonempty_file(self, tmp_path):
        fig = _make_figure()
        ax = fig.get_axes()[0]
        xs = np.linspace(0, 2 * np.pi, 5)

        def draw_frame(frame_idx: int) -> None:
            ax.clear()
            ax.plot(xs, np.sin(xs + frame_idx * 0.5))

        path = tmp_path / "test.gif"
        export_animation_gif(fig, draw_frame, n_frames=5, path=str(path), fps=5)
        assert path.exists()
        assert path.stat().st_size > 0


class TestPathTraversalValidation:
    """Issue #398: paths must be validated to prevent traversal attacks."""

    def test_rejects_empty_path(self, tmp_path):
        with pytest.raises(ValueError, match="non-empty"):
            _validate_export_path("", base_dir=tmp_path)

    def test_rejects_null_byte(self, tmp_path):
        with pytest.raises(ValueError, match="null"):
            _validate_export_path("foo\x00.png", base_dir=tmp_path)

    def test_rejects_dotdot_escape(self, tmp_path):
        with pytest.raises(ValueError, match="escapes base directory"):
            _validate_export_path("../escape.png", base_dir=tmp_path)

    def test_rejects_deep_dotdot_escape(self, tmp_path):
        with pytest.raises(ValueError, match="escapes base directory"):
            _validate_export_path("a/b/../../../escape.png", base_dir=tmp_path)

    def test_rejects_absolute_path_outside_base(self, tmp_path, tmp_path_factory):
        outside = tmp_path_factory.mktemp("outside") / "evil.png"
        with pytest.raises(ValueError, match="escapes base directory"):
            _validate_export_path(str(outside), base_dir=tmp_path)

    def test_accepts_path_inside_base(self, tmp_path):
        resolved = _validate_export_path("ok.png", base_dir=tmp_path)
        assert resolved == (tmp_path / "ok.png").resolve()

    def test_accepts_nested_subdir(self, tmp_path):
        resolved = _validate_export_path("sub/dir/ok.png", base_dir=tmp_path)
        assert resolved == (tmp_path / "sub" / "dir" / "ok.png").resolve()

    def test_accepts_absolute_path_inside_base(self, tmp_path):
        target = tmp_path / "ok.png"
        resolved = _validate_export_path(str(target), base_dir=tmp_path)
        assert resolved == target.resolve()

    def test_no_base_dir_still_rejects_null_byte(self):
        with pytest.raises(ValueError, match="null"):
            _validate_export_path("foo\x00.png", base_dir=None)

    def test_no_base_dir_still_rejects_empty(self):
        with pytest.raises(ValueError, match="non-empty"):
            _validate_export_path("", base_dir=None)

    def test_export_png_rejects_traversal(self, tmp_path):
        fig = _make_figure()
        with pytest.raises(ValueError, match="escapes base directory"):
            export_plots_png(fig, "../evil.png", base_dir=str(tmp_path))

    def test_export_pdf_rejects_traversal(self, tmp_path):
        fig = _make_figure()
        with pytest.raises(ValueError, match="escapes base directory"):
            export_plots_pdf(fig, "../evil.pdf", base_dir=str(tmp_path))

    def test_export_gif_rejects_traversal(self, tmp_path):
        fig = _make_figure()

        def draw_frame(_: int) -> None:
            pass

        with pytest.raises(ValueError, match="escapes base directory"):
            export_animation_gif(
                fig, draw_frame, n_frames=2, path="../evil.gif", base_dir=str(tmp_path)
            )

    def test_export_png_rejects_null_byte(self, tmp_path):
        fig = _make_figure()
        with pytest.raises(ValueError, match="null"):
            export_plots_png(fig, str(tmp_path / "ev\x00il.png"))

    def test_export_png_writes_inside_base(self, tmp_path):
        fig = _make_figure()
        export_plots_png(fig, "ok.png", base_dir=str(tmp_path))


class TestExportExcel:
    """Tests for export_to_excel (issue #411)."""

    @pytest.fixture(autouse=True)
    def _require_openpyxl(self):
        pytest.importorskip("openpyxl")

    def _make_result(self):
        from movement_optimizer.tests._helpers import make_test_result

        return make_test_result()

    def test_produces_nonempty_file(self, tmp_path):
        from movement_optimizer.export_excel import export_to_excel

        r = self._make_result()
        path = tmp_path / "test.xlsx"
        export_to_excel(r, str(path))
        assert path.exists()
        assert path.stat().st_size > 0

    def test_creates_parent_dirs(self, tmp_path):
        from movement_optimizer.export_excel import export_to_excel

        r = self._make_result()
        path = tmp_path / "sub" / "dir" / "test.xlsx"
        export_to_excel(r, str(path))
        assert path.exists()

    def test_has_expected_sheets(self, tmp_path):
        import openpyxl

        from movement_optimizer.export_excel import export_to_excel

        r = self._make_result()
        path = tmp_path / "test.xlsx"
        export_to_excel(r, str(path))
        wb = openpyxl.load_workbook(str(path))
        assert set(wb.sheetnames) == {"Summary", "Trajectory", "Torques", "Statistics"}

    def test_trajectory_sheet_has_correct_headers(self, tmp_path):
        import openpyxl

        from movement_optimizer.export_excel import export_to_excel

        r = self._make_result()
        path = tmp_path / "test.xlsx"
        export_to_excel(r, str(path))
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Trajectory"]
        headers = [cell.value for cell in ws[1]]
        assert headers[0] == "time (s)"
        assert "q1 (deg)" in headers
        assert "dq1 (deg/s)" in headers

    def test_torques_sheet_has_correct_headers(self, tmp_path):
        import openpyxl

        from movement_optimizer.export_excel import export_to_excel

        r = self._make_result()
        path = tmp_path / "test.xlsx"
        export_to_excel(r, str(path))
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Torques"]
        headers = [cell.value for cell in ws[1]]
        assert headers[0] == "time (s)"
        assert "tau1 (N*m)" in headers

    def test_trajectory_sheet_has_correct_row_count(self, tmp_path):
        import openpyxl

        from movement_optimizer.export_excel import export_to_excel

        r = self._make_result()
        n = len(r.t)
        path = tmp_path / "test.xlsx"
        export_to_excel(r, str(path))
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Trajectory"]
        # rows = header + n data rows
        assert ws.max_row == n + 1

    def test_torques_sheet_has_correct_row_count(self, tmp_path):
        import openpyxl

        from movement_optimizer.export_excel import export_to_excel

        r = self._make_result()
        n = len(r.t)
        path = tmp_path / "test.xlsx"
        export_to_excel(r, str(path))
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Torques"]
        assert ws.max_row == n + 1

    def test_summary_contains_exercise_name(self, tmp_path):
        import openpyxl

        from movement_optimizer.export_excel import export_to_excel

        r = self._make_result()
        path = tmp_path / "test.xlsx"
        export_to_excel(r, str(path), exercise_name="Squat")
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Summary"]
        values = [row[0].value for row in ws.iter_rows()]
        assert "Exercise" in values

    def test_summary_contains_torque_statistics(self, tmp_path):
        import openpyxl

        from movement_optimizer.export_excel import export_to_excel

        r = self._make_result()
        path = tmp_path / "test.xlsx"
        export_to_excel(r, str(path))
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Summary"]
        all_values = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value]
        assert any("Peak" in v for v in all_values)

    def test_statistics_sheet_contains_recommendations(self, tmp_path):
        import openpyxl

        from movement_optimizer.export_excel import export_to_excel

        r = self._make_result()
        path = tmp_path / "test.xlsx"
        export_to_excel(r, str(path))
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Statistics"]
        all_values = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value]
        assert "Recommendations" in all_values

    def test_raises_on_none_result(self, tmp_path):
        from movement_optimizer.export_excel import export_to_excel

        with pytest.raises(ValueError, match="None"):
            export_to_excel(None, str(tmp_path / "test.xlsx"))  # type: ignore[arg-type]

    def test_raises_on_null_byte_in_path(self, tmp_path):
        from movement_optimizer.export_excel import export_to_excel

        r = self._make_result()
        with pytest.raises(ValueError, match="null"):
            export_to_excel(r, str(tmp_path / "ev\x00il.xlsx"))
